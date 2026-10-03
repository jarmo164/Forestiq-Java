#!/usr/bin/env python3
"""Bulk import owner CSV/XLSX files into ForestIQ Postgres.

This is the large-file companion to the bounded web importer. It streams the
source into a normalized temporary CSV, uses PostgreSQL COPY into a staging
table, rejects invalid/file-duplicate rows, and inserts only owners that do not
already exist in the selected organization.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import sqlite3
import sys
import tempfile
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

FIELDS = ("displayName", "personalOrRegistryCode", "ownerType", "email", "phone", "address", "info")
REQUIRED = ("displayName", "personalOrRegistryCode")
ALIASES = {
    "displayName": {"displayname", "name", "nimi", "owner", "ownername", "omanik"},
    "personalOrRegistryCode": {"personalorregistrycode", "code", "isikukood", "registrikood", "isikuvõiregistrikood", "isikuregistrikood"},
    "ownerType": {"ownertype", "type", "tüüp", "tyyp"},
    "email": {"email", "epost"},
    "phone": {"phone", "telefon", "tel"},
    "address": {"address", "aadress"},
    "info": {"info", "notes", "note", "märkused", "markused"},
}
LIMITS = {"displayName": 300, "personalOrRegistryCode": 32, "ownerType": 64, "email": 320, "phone": 64, "address": 4000, "info": 8000}
LABELS = {"displayName": "Nimi", "personalOrRegistryCode": "Isiku- või registrikood", "ownerType": "Omaniku tüüp", "email": "E-post", "phone": "Telefon", "address": "Aadress", "info": "Info"}
CANONICAL = ("row_number", *FIELDS, "basic_valid", "validation_errors")
UUID_RE = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[1-5][0-9a-fA-F]{3}-[89abAB][0-9a-fA-F]{3}-[0-9a-fA-F]{12}$")


class ImportError(RuntimeError):
    pass


def norm(value: str) -> str:
    return "".join(ch.lower() for ch in value.strip() if ch.isalnum())


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def headers_ok(values: list[str]) -> list[str]:
    values = [text(value) for value in values]
    if not values or any(not value for value in values):
        raise ImportError("Päis puudub või sisaldab tühja veerunime.")
    if len(values) != len(set(values)):
        raise ImportError("Päises on korduvad veerunimed.")
    return values


def mapping(headers: list[str], overrides: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for header in headers:
        normalized = norm(header)
        for field, aliases in ALIASES.items():
            if field not in result and normalized in {norm(alias) for alias in aliases}:
                result[field] = header
    for item in overrides:
        if "=" not in item:
            raise ImportError(f"Vigane --map {item!r}; kasuta FIELD=HEADER.")
        field, header = (part.strip() for part in item.split("=", 1))
        if field not in FIELDS:
            raise ImportError(f"Tundmatu väli: {field}")
        if header not in headers:
            raise ImportError(f"Veergu ei leitud: {header}")
        result[field] = header
    if len(result.values()) != len(set(result.values())):
        raise ImportError("Ühte failiveergu ei saa kasutada mitme välja jaoks.")
    for field in REQUIRED:
        if field not in result:
            raise ImportError(f"Kohustuslik mapping puudub: {field}")
    return result


def validate(values: dict[str, str]) -> list[str]:
    errors: list[str] = []
    for field in REQUIRED:
        if not values[field]:
            errors.append(f"{LABELS[field]} on kohustuslik.")
    for field, maximum in LIMITS.items():
        if len(values[field]) > maximum:
            errors.append(f"{LABELS[field]} on pikem kui {maximum} märki.")
    email = values["email"]
    if email and ("@" not in email or email.startswith("@") or email.endswith("@")):
        errors.append("E-posti formaat ei ole korrektne.")
    return errors


def detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (",", ";", "\t")}
        delimiter = max(counts, key=counts.get)
        if counts[delimiter] == 0:
            raise ImportError("CSV eraldajat ei õnnestunud tuvastada.")
        return delimiter


@contextmanager
def source_rows(path: Path, encoding: str, delimiter: str | None, sheet: str | None):
    if path.suffix.lower() == ".csv":
        handle = path.open("r", encoding=encoding, newline="")
        try:
            sample = handle.read(128 * 1024)
            handle.seek(0)
            reader = csv.DictReader(handle, delimiter=delimiter or detect_delimiter(sample))
            headers = headers_ok(list(reader.fieldnames or []))
            def rows() -> Iterator[dict[str, str]]:
                for row in reader:
                    if None in row:
                        raise ImportError("CSV real on rohkem veerge kui päises.")
                    yield {header: text(row.get(header)) for header in headers}
            yield headers, rows()
        finally:
            handle.close()
        return

    if path.suffix.lower() != ".xlsx":
        raise ImportError("Toetatud on .csv ja .xlsx failid.")
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError("XLSX vajab openpyxl paketti; paigalda scripts/owners/requirements.txt.") from exc
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        ws = book[sheet] if sheet else book[book.sheetnames[0]]
        iterator = ws.iter_rows()
        try:
            first = next(iterator)
        except StopIteration as exc:
            raise ImportError("XLSX fail on tühi.") from exc
        headers = headers_ok([cell.value for cell in first])
        def rows() -> Iterator[dict[str, str]]:
            for cells in iterator:
                values = []
                for cell in cells:
                    if getattr(cell, "data_type", None) == "f":
                        raise ImportError(f"XLSX valem lahtris {getattr(cell, 'coordinate', '?')} ei ole lubatud.")
                    values.append(text(cell.value))
                values += [""] * max(0, len(headers) - len(values))
                yield dict(zip(headers, values[:len(headers)]))
        yield headers, rows()
    finally:
        book.close()


def normalize_source(args: argparse.Namespace) -> tuple[Path, int, int, int, dict[str, str]]:
    fd, name = tempfile.mkstemp(prefix="forestiq-owner-import-", suffix=".csv")
    os.close(fd)
    output = Path(name)
    total = invalid = 0
    db_fd, db_name = tempfile.mkstemp(prefix="forestiq-owner-codes-", suffix=".sqlite")
    os.close(db_fd)
    try:
        with sqlite3.connect(db_name) as codes, source_rows(args.input, args.encoding, args.delimiter, args.sheet) as (headers, rows), output.open("w", encoding="utf-8", newline="") as handle:
            codes.execute("CREATE TABLE code_count (code TEXT PRIMARY KEY, occurrences INTEGER NOT NULL)")
            selected = mapping(headers, args.map)
            writer = csv.DictWriter(handle, fieldnames=CANONICAL)
            writer.writeheader()
            for line, row in enumerate(rows, start=2):
                if args.max_rows is not None and total >= args.max_rows:
                    break
                if all(not value for value in row.values()):
                    continue
                values = {field: text(row.get(selected[field])) if field in selected else "" for field in FIELDS}
                errors = validate(values)
                writer.writerow({"row_number": line, **values, "basic_valid": not errors, "validation_errors": "; ".join(errors)})
                total += 1
                invalid += bool(errors)
                if values["personalOrRegistryCode"]:
                    codes.execute(
                        "INSERT INTO code_count(code, occurrences) VALUES (?, 1) "
                        "ON CONFLICT(code) DO UPDATE SET occurrences=occurrences+1",
                        (values["personalOrRegistryCode"],),
                    )
            duplicate_rows = int(codes.execute(
                "SELECT COALESCE(SUM(occurrences),0) FROM code_count WHERE occurrences > 1"
            ).fetchone()[0])
        if not total:
            raise ImportError("Failis ei ole andmeridu.")
        return output, total, invalid, duplicate_rows, selected
    except Exception:
        output.unlink(missing_ok=True)
        raise
    finally:
        Path(db_name).unlink(missing_ok=True)


def connect(database_url: str):
    try:
        import psycopg
    except ImportError as exc:
        raise ImportError("PostgreSQL import vajab psycopg paketti; paigalda scripts/owners/requirements.txt.") from exc
    return psycopg.connect(database_url)


def import_db(args: argparse.Namespace, normalized: Path, total: int) -> tuple[int, int, int]:
    with connect(args.database_url) as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM organization WHERE id=%s::uuid", (args.organization_id,))
            if cur.fetchone() is None:
                raise ImportError(f"Organisatsiooni ei leitud: {args.organization_id}")
            cur.execute("""
                CREATE TEMP TABLE stg_owner_bulk_import (
                    row_number bigint NOT NULL, display_name text, personal_or_registry_code text,
                    owner_type text, email text, phone text, address text, info text,
                    basic_valid boolean NOT NULL, validation_errors text
                ) ON COMMIT PRESERVE ROWS
            """)
            with cur.copy("""
                COPY stg_owner_bulk_import
                    (row_number, display_name, personal_or_registry_code, owner_type, email, phone, address, info, basic_valid, validation_errors)
                FROM STDIN WITH (FORMAT CSV, HEADER TRUE)
            """) as copy, normalized.open("rb") as handle:
                for block in iter(lambda: handle.read(1024 * 1024), b""):
                    copy.write(block)
        conn.commit()

        with conn.cursor() as cur:
            cur.execute("""
                WITH counts AS (
                    SELECT personal_or_registry_code, count(*) source_count
                    FROM stg_owner_bulk_import
                    WHERE nullif(btrim(personal_or_registry_code), '') IS NOT NULL
                    GROUP BY personal_or_registry_code
                )
                SELECT
                    count(*) FILTER (WHERE NOT s.basic_valid OR coalesce(c.source_count,0)>1),
                    count(*) FILTER (WHERE s.basic_valid AND c.source_count=1 AND o.id IS NOT NULL),
                    count(*) FILTER (WHERE s.basic_valid AND c.source_count=1 AND o.id IS NULL)
                FROM stg_owner_bulk_import s
                LEFT JOIN counts c USING (personal_or_registry_code)
                LEFT JOIN owner o ON o.organization_id=%s::uuid AND o.personal_or_registry_code=s.personal_or_registry_code
            """, (args.organization_id,))
            invalid, existing, ready = (int(value or 0) for value in cur.fetchone())

        with conn.transaction():
            with conn.cursor() as cur:
                cur.execute("""
                    WITH counts AS (
                        SELECT personal_or_registry_code, count(*) source_count
                        FROM stg_owner_bulk_import
                        WHERE nullif(btrim(personal_or_registry_code), '') IS NOT NULL
                        GROUP BY personal_or_registry_code
                    ), ready AS (
                        SELECT s.* FROM stg_owner_bulk_import s
                        JOIN counts c USING (personal_or_registry_code)
                        LEFT JOIN owner o ON o.organization_id=%s::uuid AND o.personal_or_registry_code=s.personal_or_registry_code
                        WHERE s.basic_valid AND c.source_count=1 AND o.id IS NULL
                    )
                    INSERT INTO owner (id, organization_id, display_name, personal_or_registry_code, email, phone, status, assigned_to_subject, owner_type, address, info)
                    SELECT gen_random_uuid(), %s::uuid, btrim(display_name), btrim(personal_or_registry_code),
                           nullif(btrim(email),''), nullif(btrim(phone),''), 'NEW', NULL,
                           nullif(btrim(owner_type),''), nullif(btrim(address),''), nullif(btrim(info),'')
                    FROM ready
                    ON CONFLICT (organization_id, personal_or_registry_code) DO NOTHING
                """, (args.organization_id, args.organization_id))
                inserted = max(0, cur.rowcount)
        if inserted < ready:
            existing += ready - inserted
        return inserted, existing, invalid


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("input", type=Path)
    p.add_argument("--database-url", default=os.getenv("DATABASE_URL"))
    p.add_argument("--organization-id", default=os.getenv("FORESTIQ_ORGANIZATION_ID"))
    p.add_argument("--encoding", default="utf-8-sig")
    p.add_argument("--delimiter", choices=[",", ";", "\t"])
    p.add_argument("--sheet", help="XLSX sheet; first sheet by default.")
    p.add_argument("--map", action="append", default=[], metavar="FIELD=HEADER")
    p.add_argument("--max-rows", type=int, help="Optional smoke-test limit; no default row limit.")
    p.add_argument("--dry-run", action="store_true", help="Validate/normalize without touching Postgres.")
    p.add_argument("--keep-normalized", type=Path)
    return p


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    normalized: Path | None = None
    try:
        if not args.input.is_file():
            raise ImportError(f"Sisendfaili ei leitud: {args.input}")
        if args.max_rows is not None and args.max_rows < 1:
            raise ImportError("--max-rows peab olema positiivne.")
        normalized, total, basic_invalid, duplicate_rows, selected = normalize_source(args)
        print(f"Source: {args.input}")
        print(f"SHA-256: {sha256(args.input)}")
        print(f"Rows: {total}")
        print(f"Basic invalid rows: {basic_invalid}")
        print(f"Rows affected by duplicate codes in file: {duplicate_rows}")
        print("Mapping: " + ", ".join(f"{field}={header}" for field, header in selected.items()))
        if args.keep_normalized:
            args.keep_normalized.parent.mkdir(parents=True, exist_ok=True)
            os.replace(normalized, args.keep_normalized)
            normalized = args.keep_normalized
            print(f"Normalized CSV: {normalized}")
        if args.dry_run:
            print("Dry run complete; PostgreSQL was not changed.")
            return 0
        if not args.database_url:
            raise ImportError("DATABASE_URL or --database-url is required.")
        if not args.organization_id or not UUID_RE.fullmatch(args.organization_id):
            raise ImportError("FORESTIQ_ORGANIZATION_ID or --organization-id must be a valid UUID.")
        inserted, existing, invalid = import_db(args, normalized, total)
        print(f"Inserted: {inserted}")
        print(f"Skipped existing: {existing}")
        print(f"Rejected invalid/duplicate-in-file: {invalid}")
        return 0
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    finally:
        if normalized and not args.keep_normalized:
            normalized.unlink(missing_ok=True)


if __name__ == "__main__":
    raise SystemExit(main())
